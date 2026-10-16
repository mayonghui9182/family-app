#!/usr/bin/env python3
"""以虚构文件验证解析内容、出处、告警、失败隔离与原件保护。"""
import hashlib
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

SCRIPT = Path(__file__).with_name("parse_requirements.py")
SPEC = importlib.util.spec_from_file_location("parser_under_test", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


class ParseTest(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.input = self.root / "input"
        self.input.mkdir()
        self.output = self.root / "result"

    def run_parser(self, **kwargs):
        code = MODULE.run(self.input, self.output, **kwargs)
        return code, json.loads((self.output / "extracted.json").read_text(encoding="utf-8"))

    def text(self, name, content):
        path = self.input / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
        return path

    def test_markdown_preserved_with_provenance_and_output_hashes(self):
        original = "# 规则\n\n|字段|含义|\n|---|---|\n|数量|大于0|\n\n" + chr(96) * 4 + "代码" + chr(96) * 4
        source = self.text("需求.md", original)
        original_bytes = source.read_bytes()
        code, records = self.run_parser()
        self.assertEqual(0, code)
        record = records[0]
        self.assertEqual(original, record["blocks"][0]["text"])
        self.assertEqual(original_bytes, (self.output / record["original_copy"]).read_bytes())
        self.assertEqual(original_bytes, source.read_bytes())
        manifest = json.loads((self.output / "manifest.json").read_text())
        for entry in manifest["outputs"]:
            self.assertEqual(entry["sha256"],
                hashlib.sha256((self.output / entry["path"]).read_bytes()).hexdigest())
        provenance = json.loads((self.output / "provenance.json").read_text())
        self.assertEqual(record["sha256"], provenance[0]["source_sha256"])
        self.assertIn("lines=1-", provenance[0]["locator"])

    def test_utf16_not_silently_lost(self):
        (self.input / "需求.md").write_bytes("中文必填字段".encode("utf-16"))
        code, records = self.run_parser()
        self.assertEqual(1, code)
        self.assertEqual("中文必填字段", records[0]["blocks"][0]["text"])
        self.assertEqual("utf-16", records[0]["inventory"]["encoding"])

    def test_word_tables_headers_and_revisions(self):
        path = self.input / "需求.docx"
        doc = Document()
        doc.add_heading("审核规则", 1)
        doc.add_paragraph("只有管理员可删除")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "字段说明"
        table.cell(1, 0).text = "金额"
        table.cell(1, 1).text = "不得为负"
        doc.sections[0].header.paragraphs[0].text = "内部需求v1"
        doc.save(path)
        with zipfile.ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        original = parts["word/document.xml"].decode()
        original = original.replace("</w:body>",
            '<w:p><w:del w:id="1"><w:r><w:delText>旧规则</w:delText></w:r></w:del>'
            '<w:ins w:id="2"><w:r><w:t>新规则</w:t></w:r></w:ins></w:p></w:body>')
        parts["word/document.xml"] = original.encode()
        parts["word/embeddings/object.bin"] = b"fake embedded object"
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
        code, records = self.run_parser()
        self.assertEqual(1, code)
        record = records[0]
        texts = "\n".join(x["text"] for x in record["blocks"])
        for expected in ("只有管理员可删除", "字段说明", "不得为负", "内部需求v1", "旧规则", "新规则"):
            self.assertIn(expected, texts)
        part = record["inventory"]["word_parts"]["word/document.xml"]
        self.assertEqual(part["source_text_chars"], part["extracted_text_chars"])
        self.assertEqual(2, len(part["tables"][0]["rows"]))
        self.assertIn("gridSpan", part["tables"][0]["rows"][0]["cells"][0]["properties_xml"])
        self.assertTrue(any(x["code"] == "embedded_object" for x in record["issues"]))
        self.assertTrue(any(x["code"] == "word_special_content" for x in record["issues"]))

    def test_excel_values_formula_styles_hidden_and_coordinates(self):
        path = self.input / "字段.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "业务字段"
        sheet.append(["字段", "取值"])
        sheet["A2"] = "金额"
        sheet["B2"] = 0
        sheet["B3"] = False
        sheet["C2"] = "=B2+1"
        sheet["A2"].fill = PatternFill("solid", fgColor="FF0000")
        sheet["A2"].comment = Comment("红色为必填", "需求负责人")
        sheet.row_dimensions[2].hidden = True
        sheet.merge_cells("A4:B4")
        sheet["A4"] = "合并标题"
        hidden = book.create_sheet("枚举")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "拒绝"
        book.save(path)
        code, records = self.run_parser()
        self.assertEqual(1, code, records[0]["issues"])
        record = records[0]
        cells = {x["locator"]: x for x in record["blocks"]}
        self.assertEqual(0, cells["sheet=业务字段 / cell=B2"]["metadata"]["value"])
        self.assertIs(False, cells["sheet=业务字段 / cell=B3"]["metadata"]["value"])
        formula = cells["sheet=业务字段 / cell=C2"]["metadata"]
        self.assertEqual("=B2+1", formula["formula"])
        self.assertIsNone(formula["cached_result"])
        self.assertEqual("红色为必填", cells["sheet=业务字段 / cell=A2"]["metadata"]["comment"]["text"])
        self.assertIn("FF0000", cells["sheet=业务字段 / cell=A2"]["metadata"]["style"]["fill"])
        self.assertIn("sheet=枚举 / cell=A1", cells)
        self.assertTrue(any(x["code"] == "formula_not_calculated" for x in record["issues"]))
        self.assertIn("A4:B4", record["inventory"]["sheets"][0]["merged_ranges"])

    def test_html_static_controls_local_and_remote_images(self):
        self.text("原型/index.html",
            '<html><style>.hide{display:none}</style><body><h1>审核</h1>'
            '<input id="name" required><button disabled>提交</button>'
            '<div hidden>错误提示</div><img src="image.png">'
            '<img src="https://example.invalid/remote.png">'
            '<script>throw new Error("不得执行")</script></body></html>')
        (self.input / "原型/image.png").write_bytes(b"fake-png")
        code, records = self.run_parser()
        self.assertEqual(1, code)
        record = next(r for r in records if r["format"] == ".html")
        self.assertEqual(2, len(record["inventory"]["controls"]))
        self.assertIn("required", record["inventory"]["controls"][0]["attributes"])
        texts = "".join(x["text"] for x in record["blocks"])
        self.assertIn("错误提示", texts)
        self.assertNotIn("不得执行", texts)
        self.assertEqual(1, len(record["assets"]))
        self.assertEqual(b"fake-png", (self.output / record["assets"][0]["path"]).read_bytes())
        self.assertTrue(any(x["code"] == "html_image_not_copied" for x in record["issues"]))

    def test_html_outside_image_not_read(self):
        (self.root / "private.png").write_bytes(b"private")
        self.text("页面.html", '<img src="../private.png">')
        _, records = self.run_parser()
        self.assertEqual([], records[0]["assets"])

    def test_image_is_pending_ocr(self):
        (self.input / "页面.png").write_bytes(b"image-payload")
        code, records = self.run_parser()
        self.assertEqual(1, code)
        self.assertEqual("review_required", records[0]["status"])
        self.assertEqual("ocr_pending", records[0]["issues"][0]["code"])

    def test_unsupported_and_corrupt_isolate_other_results(self):
        self.text("正确.md", "有效需求")
        (self.input / "损坏.docx").write_bytes(b"not-a-zip")
        (self.input / "旧文件.doc").write_bytes(b"old-format")
        code, records = self.run_parser()
        self.assertEqual(2, code)
        self.assertEqual({"parsed", "failed", "unsupported"}, {r["status"] for r in records})
        self.assertTrue((self.output / "report.md").exists())

    def test_existing_output_not_overwritten(self):
        self.text("需求.md", "需求")
        self.output.mkdir()
        marker = self.output / "marker"
        marker.write_text("保留")
        with self.assertRaises(ValueError):
            MODULE.run(self.input, self.output)
        self.assertEqual("保留", marker.read_text())

    def test_nested_output_rejected_to_avoid_future_recursion(self):
        self.text("需求.md", "需求")
        self.output = self.input / "解析结果"
        with self.assertRaises(ValueError):
            self.run_parser()

    def test_cell_limit_report_failure(self):
        path = self.input / "大表.xlsx"
        book = Workbook()
        book.active["Z100"] = "远处的内容"
        book.save(path)
        code, records = self.run_parser(max_cells=10)
        self.assertEqual(2, code)
        self.assertIn("--max-cells", records[0]["issues"][-1]["message"])

    def test_file_limit(self):
        (self.input / "巨大.md").write_bytes(b"x" * (1024 * 1024 + 1))
        code, records = self.run_parser(max_file_mb=1)
        self.assertEqual(2, code)
        self.assertNotIn("original_copy", records[0])

    def test_package_bomb_rejected_before_library(self):
        path = self.input / "展开过大.docx"
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("large.xml", "x" * (1024 * 1024 + 1))
        code, records = self.run_parser(max_expanded_mb=1)
        self.assertEqual(2, code)
        self.assertIn("解压大小", records[0]["issues"][-1]["message"])

    def test_missing_dependency_is_explicit(self):
        doc = Document()
        doc.add_paragraph("需求")
        doc.save(self.input / "需求.docx")
        with patch.object(MODULE, "parse_docx", side_effect=ModuleNotFoundError("No module named docx")):
            code, records = self.run_parser()
        self.assertEqual(2, code)
        self.assertIn("ModuleNotFoundError", records[0]["issues"][-1]["message"])

    def test_merged_cell_loss_is_reported_against_raw_xml(self):
        path = self.input / "合并表.xlsx"
        book = Workbook()
        book.active["A1"] = "主单元格"
        book.active.merge_cells("A1:B1")
        book.save(path)
        with zipfile.ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        xml = parts["xl/worksheets/sheet1.xml"].decode()
        xml = xml.replace("</row>", '<c r="B1" t="inlineStr"><is><t>原始被覆盖文字</t></is></c></row>')
        parts["xl/worksheets/sheet1.xml"] = xml.encode()
        with zipfile.ZipFile(path, "w") as archive:
            for name, value in parts.items():
                archive.writestr(name, value)
        code, records = self.run_parser()
        self.assertEqual(1, code)
        self.assertIn("B1", records[0]["inventory"]["sheets"][0]["source_cells_not_emitted"])
        self.assertTrue(any(x["code"] == "excel_cell_coverage" for x in records[0]["issues"]))

    def test_library_warning_is_not_silent(self):
        import warnings
        self.text("需求.md", "规则")
        original = MODULE.parse_md
        def warning_parser(*args):
            warnings.warn("虚构的解析能力缺失")
            return original(*args)
        with patch.object(MODULE, "parse_md", side_effect=warning_parser):
            code, records = self.run_parser()
        self.assertEqual(1, code)
        self.assertTrue(any(x["code"] == "library_warning" for x in records[0]["issues"]))

    def test_cli_exit_codes(self):
        self.text("需求.md", "明确需求")
        result = subprocess.run([sys.executable, str(SCRIPT), str(self.input),
                                 "--output", str(self.output)], capture_output=True)
        self.assertEqual(0, result.returncode, result.stderr)
        result = subprocess.run([sys.executable, str(SCRIPT), str(self.input),
                                 "--output", str(self.output)], capture_output=True)
        self.assertEqual(2, result.returncode)

    @unittest.skipIf(sys.platform == "win32", "Windows普通权限未必允许创建符号链接")
    def test_symlink_not_followed(self):
        secret = self.root / "other.md"
        secret.write_text("不应读入")
        (self.input / "链接.md").symlink_to(secret)
        code, records = self.run_parser()
        self.assertEqual(2, code)
        self.assertEqual([], records[0]["blocks"])


if __name__ == "__main__":
    unittest.main()
